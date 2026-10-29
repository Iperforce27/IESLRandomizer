from tkinter import *
import openpyxl
import random
import zipfile
from io import BytesIO
from PIL import Image, ImageTk

#Class that represents the zip file for images
class Archivio:
    def __init__(self, path="Images.zip"):
        self.z = zipfile.ZipFile(path)          # aperto una volta sola
        self.nomi = set(self.z.namelist())

    def get(self, nome):
        return Image.open(BytesIO(self.z.read(nome)))

    def __len__(self):
        return len(self.nomi)

    def close(self):
        self.z.close()

#Class that represents a player or staff
class OC:
    def __init__(self, name, team, image, arch):
        self.name = name    #The player's name
        self.team = team    #The player's team
        if(image not in ['N/A', '#VALUE!']):
            self.image = ImageTk.PhotoImage(arch.get(f'{image}.webp').resize([50,50]))  #The player's image
        else:
            self.image = image  #No image
        pass

excel_file = None   #where the excel file is stored

playerList = [] #List of players extracted
staffList = []  #List of coach and manager extracted
#DFnumb = None   #Number of defenders to extract
#MFnumb = None   #Number of midfielders to extract
#FWnumb = None   #Number of forwards to extract
rand_staff = None  #Tells whether to randomize the staff

#Check whether the player is already in the team
def CheckDuplicate(player):
    alreadyIn = []
    for p in playerList:
        alreadyIn.append(p.name)
    return (player.name in alreadyIn)

#Add a random player to the team
def GetRandomPlayer(sheet, arch):
    index = random.randrange(1, sheet.max_row)
    p = OC(sheet[f"D{index}"].value, sheet[f"C{index}"].value, sheet[f"E{index}"].value, arch)
    while CheckDuplicate(p):
        index = random.randrange(1, sheet.max_row)
        p = OC(sheet[f"D{index}"].value, sheet[f"C{index}"].value, sheet[f"E{index}"].value, arch)
    playerList.append(p)

#Add a random staff member
def GetRandomStaff(sheet, arch):
    index = random.randrange(1,sheet.max_row)
    s = OC(sheet[f"D{index}"].value, sheet[f"C{index}"].value, sheet[f"E{index}"].value, arch)
    staffList.append(s)

#Randomize the team
def Randomize(DFn, MFn, FWn):
    a = Archivio()
    excel_file = openpyxl.load_workbook("IESL.xlsx")
    GetRandomPlayer(excel_file['GKs'], a)
    for i in range(DFn):
        GetRandomPlayer(excel_file['DFs'], a)
    for i in range(MFn):
        GetRandomPlayer(excel_file['MFs'], a)
    for i in range(FWn):
        GetRandomPlayer(excel_file['FWs'], a)
    if rand_staff:
        GetRandomStaff(excel_file['COs'], a)
        GetRandomStaff(excel_file['MGs'], a)
    excel_file.close()
    a.close()

#Show the randomized team in a new window
def PrintTeam():
    teamPage = Toplevel(root)
    teamPage.title("Your team")
    Label(teamPage, text="Players:").grid(row=0, column=0)
    i = 1
    for player in playerList:
        Label(teamPage, text=player.name).grid(row=i,column=0)
        Label(teamPage, text=f'({player.team})').grid(row=i,column=1)
        if(player.image not in ['N/A', '#VALUE!']):
            Label(teamPage, image=player.image).grid(row=i,column=2)
        i=i+1
    if(rand_staff):
        Label(teamPage, text="Staff:").grid(row=i, column=0)
        i=i+1
        for staff in staffList:
            Label(teamPage, text=staff.name).grid(row=i,column=0)
            Label(teamPage, text=f'({staff.team})').grid(row=i,column=1)
            if(staff.image not in ['N/A', '#VALUE!']):
                Label(teamPage, image=staff.image).grid(row=i,column=2)
            i=i+1
    teamPage.mainloop()

#Design window
root = Tk()
root.title("IESL Randomizer")
root.geometry("600x400")
root.iconphoto(True, PhotoImage("Logo", file="logo.png"))

titleLabel = Label(root, text="IESL Randomizer (made by Iper)", height=2).pack()
introLabel = Label(root, text="Welcome to the IESL Randomizer!").pack(anchor='w')
instructionLabel = Label(root, text="Input the formation").pack(anchor='w')

inputFrame = Frame(root)
Label(inputFrame, text="Defenders:").grid(row=0,column=0)
Label(inputFrame, text="Midfielders:").grid(row=1,column=0)
Label(inputFrame, text="Forwards:").grid(row=2,column=0)
Label(inputFrame, text="Randomize staff").grid(row=3,column=0)
DFnumb_var = IntVar(inputFrame)
MFnumb_var = IntVar(inputFrame)
FWnumb_var = IntVar(inputFrame)
staff_var = BooleanVar(inputFrame)
Entry(inputFrame, textvariable=DFnumb_var).grid(row=0, column=1)
Entry(inputFrame, textvariable=MFnumb_var).grid(row=1, column=1)
Entry(inputFrame, textvariable=FWnumb_var).grid(row=2, column=1)
Checkbutton(inputFrame, variable=staff_var).grid(row=3, column=1)
inputFrame.pack(anchor='w')

#Start the randomization
def StartRandomize():
    playerList.clear()
    staffList.clear()
    #DFnumb = DFnumb_var.get()
    #MFnumb = MFnumb_var.get()
    #FWnumb = FWnumb_var.get()
    global rand_staff
    rand_staff = staff_var.get()
    Randomize(DFnumb_var.get(), MFnumb_var.get(), FWnumb_var.get())
    PrintTeam()
    
Button(root, text="Randomize!", background="Orange", command=StartRandomize).pack(anchor="w")
root.mainloop()
